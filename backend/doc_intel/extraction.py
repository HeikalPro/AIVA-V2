"""document-extractor integration: the API-process side of the extraction stage (Flow 1).

The same child-process machinery also runs Flow 2's CRM extraction (``crm_extraction.py``):
both share ``_EXTRACTION_SLOT`` (one extraction child per process, across both flows), the
active-children registry that ``terminate_active_extractions`` kills on shutdown, the
hard timeout and the scrubbed child environment.

Extraction always runs in a CHILD PROCESS: document-extractor's native engines
(PDFium, Tesseract) are not thread-safe, OCR is CPU heavy, and a crashing parser
must never take the API process (which also serves chat) down with it.

``run_extraction``:
- writes a job file;
- starts ``python -m backend.doc_intel.extraction_worker`` with a hard timeout. The
  whole process tree, Tesseract included, is killed when the timeout passes;
- maps the outcome to a NormalizedDocument, or to an ExtractionFailed whose reason
  an admin can act on.

The child runs at lower CPU priority, uses one OCR thread, and gets a scrubbed
environment:
- no ``DOCUMENT_EXTRACTOR_CONFIG``, so the repo's document_extractor.toml (which turns
  LLM "intelligence" on) is never loaded;
- no API keys, passwords or tokens.
Every extractor option is explicit, with intelligence off, so nothing leaves the server.

``ExtractionFailed.code`` values:
- ``unavailable``, ``source_missing``, ``unsupported``, ``timeout``, ``crash``;
- from the child: ``password_protected``, ``resource_limit``, ``too_many_pages``,
  ``parse_error``, ``missing_dependency``, ``ocr_failed``, ``config_error``,
  ``extraction_failed``, ``no_text``.
"""
from __future__ import annotations

import hashlib
import importlib
import importlib.metadata
import importlib.util
import json
import logging
import os
import re
import shutil
import signal
import subprocess
import sys
import tempfile
import threading
import time
from pathlib import Path
from typing import Any

from backend.doc_intel.constants import ALLOWED_EXTENSIONS
from backend.doc_intel.extraction_worker import JOB_VERSION, scrub_detail
from backend.doc_intel.normalized import NormalizedDocument
from backend.doc_intel.settings import DocIntelSettings

_log = logging.getLogger(__name__)

# AIVA-V2/: the child's working directory, so ``-m backend.doc_intel.extraction_worker`` resolves.
_ROOT = Path(__file__).resolve().parent.parent.parent
_WORKER_MODULE = "backend.doc_intel.extraction_worker"

PAGE_OCR_DPI = 300  # page rendering for the Arabic OCR fallback
_REAP_SECONDS = 10.0  # after a timeout kill, how long to wait for the process tree to go
_PROBE_TIMEOUT_SECONDS = 5.0  # tesseract --version / --list-langs
_WINDOWS_TESSERACT = r"C:\Program Files\Tesseract-OCR\tesseract.exe"

_PDF = ALLOWED_EXTENSIONS[".pdf"]
_DOCX = ALLOWED_EXTENSIONS[".docx"]
_KIND_BY_MEDIA_TYPE = {_PDF: "pdf", _DOCX: "docx"}

# Imported lazily by document-extractor; a missing one only shows up mid-extraction.
_RUNTIME_MODULES = (("pypdfium2", "pypdfium2"), ("docx", "python-docx"), ("lxml", "lxml"), ("PIL", "Pillow"))

_INSTALL_ACTION = (
    "Install document-extractor from the wheel and backend/requirements-docintel.txt in the backend "
    "environment (see the runbook), then restart the backend"
)
_TIMEOUT_ACTION = (
    "Split the document into smaller files, or retry when the server is less busy "
    "(the limit is DOC_INTEL_EXTRACTION_TIMEOUT_SECONDS)"
)

# The child gets the parent's environment minus anything that looks like a secret.
_SECRET_NAME = re.compile(
    r"KEY|SECRET|PASSWORD|PASSWD|PASSPHRASE|TOKEN|CREDENTIAL|PRIVATE|COOKIE|SESSION|AUTH|DSN|CONNECTION_STRING",
    re.IGNORECASE,
)
_SECRET_PREFIXES = (
    "DOCUMENT_EXTRACTOR_", "DOC_INTEL_", "ORACLE_", "OPENAI_", "SOVEREIGNEG_", "ANTHROPIC_", "AZURE_",
    "LLAMA_", "SMTP_", "ZOHO_", "JWT_", "BOOTSTRAP_", "AWS_", "GOOGLE_", "REDIS", "DATABASE_",
    # crm-document-ingestion's own settings (Microsoft credentials, CRM options): the CRM job
    # passes every option explicitly, so none of these may reach the child.
    "MICROSOFT_", "CRM_",
)
_URL_CREDENTIALS = re.compile(r"://[^/\s:@]*:[^/\s@]*@")

_EXTRACTION_SLOT = threading.BoundedSemaphore(1)
# Running extraction children, so shutdown can kill them (terminate_active_extractions).
_ACTIVE_CHILDREN: set[subprocess.Popen[bytes]] = set()
_ACTIVE_LOCK = threading.Lock()

_SMOKE_MARKER = "AIVA-SMOKE-7F3C9"
_SMOKE_ARABIC = "هذا اختبار للتأكد من أن استخراج النص العربي يعمل"


class ExtractionFailed(Exception):
    """Extraction could not produce usable text; ``reason`` is shown to the admin as-is."""

    def __init__(self, reason: str, *, code: str = "extraction_failed", suggested_action: str | None = None) -> None:
        super().__init__(reason)
        self.reason = reason
        self.code = code
        self.suggested_action = suggested_action


def extraction_available() -> tuple[bool, str | None]:
    """(True, None) when document-extractor can be imported in this environment,
    else (False, user-facing reason). Cheap; safe to call per request."""
    try:
        # Pure Python at import time (~0.1 s once, then cached): no native engine is loaded.
        importlib.import_module("document_extractor")
    except ImportError:
        return False, "document-extractor is not installed in the backend environment — see the runbook"
    except Exception as exc:  # a broken install
        return False, f"document-extractor cannot be imported ({type(exc).__name__})"
    for module, package in _RUNTIME_MODULES:
        try:
            missing = importlib.util.find_spec(module) is None
        except (ImportError, ValueError):
            missing = True
        if missing:
            return False, f"The server is missing an extraction component: {package}"
    return True, None


def extractor_info(settings: DocIntelSettings) -> dict[str, Any]:
    """Versions and OCR capabilities for diagnostics: document_extractor version,
    tesseract path/version, installed vs requested OCR languages. Never raises."""
    requested = list(settings.ocr_language_list)
    info: dict[str, Any] = {
        "available": False,
        "reason": None,
        "document_extractor_version": None,
        "pypdfium2_version": None,
        "pdf_engine": "pdfium",
        "mode": settings.extraction_mode,
        "ocr_languages_requested": requested,
        # None = unknown (Tesseract missing, or its languages could not be listed).
        "ocr_languages_installed": None,
        "ocr_languages_missing": list(requested),
        "pytesseract_installed": False,
        # The binary was found AND runs (its version could be read).
        "tesseract_available": False,
        # Tesseract available, pytesseract installed, and every requested language present.
        "ocr_available": False,
        "tesseract": {
            "available": False, "found": False, "path": None, "configured": None, "version": None, "error": None,
        },
    }
    try:
        info["available"], info["reason"] = extraction_available()
        info["document_extractor_version"] = _package_version("document-extractor")
        info["pypdfium2_version"] = _package_version("pypdfium2")
        info["pytesseract_installed"] = importlib.util.find_spec("pytesseract") is not None

        env = _child_env(settings)
        tesseract = info["tesseract"]
        path, configured = _resolve_tesseract(settings, env)
        tesseract["configured"] = configured
        if path is None:
            tesseract["error"] = (
                "The configured Tesseract binary was not found (DOC_INTEL_TESSERACT_CMD / TESSERACT_CMD)"
                if configured else "Tesseract is not installed or not on PATH"
            )
        else:
            tesseract["found"] = True
            tesseract["path"] = path
            tesseract["version"], version_error = _tesseract_version(path, env)
            tesseract["available"] = info["tesseract_available"] = tesseract["version"] is not None
            installed, langs_error = _tesseract_languages(path, env)
            tesseract["error"] = version_error or langs_error
            if langs_error is None:
                info["ocr_languages_installed"] = installed
                info["ocr_languages_missing"] = [lang for lang in requested if lang not in installed]
        info["ocr_available"] = bool(
            info["tesseract_available"] and info["pytesseract_installed"]
            and info["ocr_languages_installed"] is not None and not info["ocr_languages_missing"]
        )
    except Exception as exc:  # diagnostics must never break the caller
        _log.warning("doc_intel: extractor_info failed", exc_info=True)
        info["error"] = f"Diagnostics failed ({type(exc).__name__})"
    return info


def run_extraction(
    source_path: Path,
    *,
    filename: str,
    media_type: str,
    sha256: str,
    out_path: Path,
    settings: DocIntelSettings,
) -> NormalizedDocument:
    """BLOCKING (call via asyncio.to_thread). Extract ``source_path`` in a child process
    with a hard timeout (settings.extraction_timeout_seconds), write the normalized
    document to ``out_path`` and return it.

    Raises ExtractionFailed with a user-facing reason (password-protected, corrupt,
    timeout, no text / OCR unavailable, page limit exceeded, ...).
    """
    # One extraction child per process at a time: OCR is CPU heavy and the API
    # process also serves chat (plan risk R2).
    with _EXTRACTION_SLOT:
        return _run_child(
            Path(source_path), filename=filename, media_type=media_type, sha256=sha256,
            out_path=Path(out_path), settings=settings, timeout=float(settings.extraction_timeout_seconds),
        )


def run_smoke_test(settings: DocIntelSettings, *, timeout_seconds: float) -> dict[str, Any]:
    """BLOCKING. Extract a tiny generated DOCX in a child process (health check).

    Returns {"ok": bool, "seconds": float, "detail": str | None, "info": extractor_info(...)}.
    When a document is being extracted right now the smoke test is skipped rather than
    queued behind it (a long OCR job would otherwise time the health check out), and
    the result carries ``"busy": True``. Never raises.
    """
    if not _EXTRACTION_SLOT.acquire(blocking=False):
        return {
            "ok": True,
            "busy": True,
            "seconds": 0.0,
            "detail": "A document is being extracted right now, so the smoke test was skipped",
            "info": extractor_info(settings),
        }
    try:
        return _smoke_test_unlocked(settings, timeout_seconds=timeout_seconds)
    finally:
        _EXTRACTION_SLOT.release()


def _smoke_test_unlocked(settings: DocIntelSettings, *, timeout_seconds: float) -> dict[str, Any]:
    started = time.perf_counter()
    ok = False
    detail: str | None = None
    try:
        with tempfile.TemporaryDirectory(prefix="aiva-docintel-smoke-", ignore_cleanup_errors=True) as tmp:
            source = Path(tmp) / "smoke.docx"
            _write_smoke_docx(source)
            document = _run_child(
                source, filename="smoke-test.docx", media_type=_DOCX,
                sha256=hashlib.sha256(source.read_bytes()).hexdigest(),
                out_path=Path(tmp) / "normalized.json", settings=settings, timeout=float(timeout_seconds),
            )
            text = "\n".join(b.text for b in document.blocks)
            if _SMOKE_MARKER not in text:
                detail = "The smoke document was extracted, but its marker text is missing"
            elif _SMOKE_ARABIC not in text:
                detail = "The smoke document was extracted, but its Arabic text did not come through intact"
            else:
                ok = True
    except ExtractionFailed as exc:
        if exc.code == "timeout":
            detail = f"Smoke test timed out after {float(timeout_seconds):g} s"
        else:
            detail = f"Smoke test failed: {exc.reason}"
    except ImportError as exc:
        detail = f"Smoke test needs python-docx, which is not installed ({exc.name or 'docx'})"
    except Exception as exc:
        _log.warning("doc_intel: extraction smoke test failed", exc_info=True)
        detail = f"Smoke test failed unexpectedly ({type(exc).__name__})"
    seconds = round(time.perf_counter() - started, 2)
    return {"ok": ok, "seconds": seconds, "detail": detail, "info": extractor_info(settings)}


# ---- the child process ---------------------------------------------------------------------


def _run_child(
    source_path: Path,
    *,
    filename: str,
    media_type: str,
    sha256: str,
    out_path: Path,
    settings: DocIntelSettings,
    timeout: float,
) -> NormalizedDocument:
    available, why = extraction_available()
    if not available:
        raise ExtractionFailed(why or "document-extractor is not available", code="unavailable",
                               suggested_action=_INSTALL_ACTION)
    if not source_path.is_file():
        raise ExtractionFailed("The stored file is missing — upload the document again", code="source_missing")
    media_type = _media_type_for(media_type, source_path)

    # On failure out_path must not exist (a stale result would look like a finished extraction).
    _remove_quietly(out_path)
    _remove_quietly(_tmp_sibling(out_path))
    job_dir = Path(tempfile.mkdtemp(prefix="aiva-docintel-"))
    try:
        error_path = job_dir / "error.json"
        job = _job_base(source_path, filename=filename, media_type=media_type, sha256=sha256,
                        settings=settings, timeout=timeout)
        job.update({"output_path": str(out_path.resolve()), "error_path": str(error_path)})
        started = time.perf_counter()
        try:
            result = _spawn_job(job, job_dir / "job.json", settings, timeout)
        except subprocess.TimeoutExpired:
            _log.warning("doc_intel: extraction of %s killed after %.0f s", sha256[:12], timeout)
            raise ExtractionFailed(_timeout_reason(timeout), code="timeout", suggested_action=_TIMEOUT_ACTION) from None
        elapsed = time.perf_counter() - started

        if result.returncode == 0:
            try:
                document = NormalizedDocument.load(out_path)
            except (OSError, ValueError) as exc:
                raise ExtractionFailed("The extraction process finished without a usable result",
                                       code="crash") from exc
            _log.info(
                "doc_intel: extracted %s in %.1f s: %s pages, %s blocks, %s chars%s",
                sha256[:12], elapsed, document.page_count, len(document.blocks), document.text_chars,
                " (Arabic page-OCR fallback)" if document.extractor.get("fallback") else "",
            )
            return document

        failure = _read_failure(error_path, timeout)
        if failure is None:
            failure = ExtractionFailed(
                _crash_reason(result), code="crash",
                suggested_action="Retry; if it fails again, check the server log for the extraction error",
            )
        _log.warning(
            "doc_intel: extraction of %s failed after %.1f s (%s, exit %s): %s",
            sha256[:12], elapsed, failure.code, result.returncode, _stderr_tail(result.stderr) or "-",
        )
        raise failure
    except BaseException:
        _remove_quietly(out_path)
        raise
    finally:
        _remove_quietly(_tmp_sibling(out_path))
        shutil.rmtree(job_dir, ignore_errors=True)


# ---- shared with crm_extraction (package-internal) -----------------------------------------


def _job_base(
    source_path: Path,
    *,
    filename: str,
    media_type: str,
    sha256: str,
    settings: DocIntelSettings,
    timeout: float,
) -> dict[str, Any]:
    """The explicit extraction options every worker job carries (Flow 1 and CRM jobs alike)."""
    return {
        "version": JOB_VERSION,
        "input_path": str(source_path.resolve()),
        "filename": filename,
        "media_type": media_type,
        "sha256": sha256,
        "mode": settings.extraction_mode.strip().lower(),
        "ocr_languages": list(settings.ocr_language_list),
        "max_pages": settings.max_pages,
        "tesseract_cmd": (settings.tesseract_cmd or "").strip() or None,
        "page_ocr_dpi": PAGE_OCR_DPI,
        "timeout_seconds": timeout,
    }


def _spawn_job(
    job: dict[str, Any], job_path: Path, settings: DocIntelSettings, timeout: float
) -> subprocess.CompletedProcess[bytes]:
    """Write ``job`` to ``job_path`` and run the worker on it: scrubbed environment, lower
    priority, process-tree kill at ``timeout`` (raises subprocess.TimeoutExpired)."""
    job_path.write_text(json.dumps(job, ensure_ascii=False), encoding="utf-8")
    command = [sys.executable, "-m", _WORKER_MODULE, "--job", str(job_path)]
    return _spawn(command, cwd=_ROOT, env=_child_env(settings), timeout=timeout)


def _read_error_payload(error_path: Path) -> dict[str, Any] | None:
    """The worker's error object ({"code", "reason", "suggested_action"[, "stage"]}), or None."""
    try:
        payload = json.loads(error_path.read_text(encoding="utf-8"))
        error = payload["error"]
    except (OSError, ValueError, KeyError, TypeError):
        return None
    return error if isinstance(error, dict) else None


def _media_type_for(media_type: str, source_path: Path) -> str:
    if media_type in _KIND_BY_MEDIA_TYPE:
        return media_type
    by_suffix = ALLOWED_EXTENSIONS.get(source_path.suffix.lower())
    if by_suffix is None:
        raise ExtractionFailed("Unsupported or corrupt file", code="unsupported")
    return by_suffix


def _child_env(settings: DocIntelSettings) -> dict[str, str]:
    """The parent's environment without DOCUMENT_EXTRACTOR_CONFIG or anything secret-looking.

    The child needs no secret. In production os.environ does hold some (compose
    env_file, and zoho_auth loads its .env into os.environ), so they are dropped by
    name (KEY, SECRET, PASSWORD, TOKEN, ... and known prefixes) and by value
    (URLs with embedded credentials).
    """
    env: dict[str, str] = {}
    for name, value in os.environ.items():
        upper = name.upper()
        if upper.startswith(_SECRET_PREFIXES) or _SECRET_NAME.search(upper) or _URL_CREDENTIALS.search(value):
            continue
        env[name] = value
    tesseract_cmd = (settings.tesseract_cmd or "").strip()
    if tesseract_cmd:
        env["TESSERACT_CMD"] = tesseract_cmd
    env.setdefault("OMP_THREAD_LIMIT", "1")  # one Tesseract thread: chat shares this machine
    env["PYTHONIOENCODING"] = "utf-8"
    existing = env.get("PYTHONPATH")
    env["PYTHONPATH"] = str(_ROOT) + (os.pathsep + existing if existing else "")
    return env


def _spawn(command: list[str], *, cwd: Path, env: dict[str, str], timeout: float) -> subprocess.CompletedProcess[bytes]:
    """subprocess.run with a hard timeout that kills the whole process tree.

    subprocess.run's own timeout kills only the direct child. That would leave
    Tesseract running, and on Windows the real interpreter behind a venv's
    python.exe launcher.
    """
    kwargs: dict[str, Any] = {}
    if os.name == "nt":
        kwargs["creationflags"] = subprocess.CREATE_NEW_PROCESS_GROUP | subprocess.BELOW_NORMAL_PRIORITY_CLASS
    else:
        kwargs["start_new_session"] = True  # own process group, killed as one; the worker lowers its priority
    proc = subprocess.Popen(  # noqa: S603 - fixed argv built here, never from document content
        command, cwd=str(cwd), env=env, stdin=subprocess.DEVNULL, stdout=subprocess.PIPE, stderr=subprocess.PIPE,
        **kwargs,
    )
    with _ACTIVE_LOCK:
        _ACTIVE_CHILDREN.add(proc)
    try:
        stdout, stderr = proc.communicate(timeout=timeout)
    except subprocess.TimeoutExpired:
        _kill_tree(proc)
        try:
            proc.communicate(timeout=_REAP_SECONDS)
        except Exception:
            pass
        raise
    except BaseException:
        _kill_tree(proc)
        raise
    finally:
        with _ACTIVE_LOCK:
            _ACTIVE_CHILDREN.discard(proc)
    return subprocess.CompletedProcess(command, proc.returncode, stdout, stderr)


def terminate_active_extractions() -> int:
    """Kill every running extraction child (and its OCR processes); used on shutdown.

    Without this a restart waits for an in-flight OCR job (up to the extraction
    timeout), because the waiting thread keeps the interpreter alive, and a forced
    exit leaves the child running as an orphan. Returns how many were killed.
    """
    with _ACTIVE_LOCK:
        children = list(_ACTIVE_CHILDREN)
    for proc in children:
        _kill_tree(proc)
    return len(children)


def _kill_tree(proc: subprocess.Popen[bytes]) -> None:
    if os.name == "nt":
        try:
            subprocess.run(
                ["taskkill", "/F", "/T", "/PID", str(proc.pid)],
                capture_output=True, timeout=_REAP_SECONDS, check=False,
            )
        except (OSError, subprocess.SubprocessError):
            pass
    else:
        try:
            os.killpg(proc.pid, signal.SIGKILL)  # start_new_session: the group id is the child's pid
        except (ProcessLookupError, PermissionError, OSError):
            pass
    try:
        proc.kill()
    except OSError:
        pass


def _read_failure(error_path: Path, timeout: float) -> ExtractionFailed | None:
    error = _read_error_payload(error_path)
    if error is None:
        return None
    code = str(error.get("code") or "extraction_failed")
    reason = str(error.get("reason") or "").strip()
    action = error.get("suggested_action")
    if code == "timeout":
        return ExtractionFailed(_timeout_reason(timeout), code="timeout",
                                suggested_action=scrub_detail(action, 400) if action else _TIMEOUT_ACTION)
    if not reason:
        return None
    return ExtractionFailed(scrub_detail(reason, 400), code=code,
                            suggested_action=scrub_detail(action, 400) if action else None)


def _timeout_reason(seconds: float) -> str:
    return f"Extraction timed out after {float(seconds):g} s — the file may be very large or scanned; split it or retry"


def _crash_reason(result: subprocess.CompletedProcess[bytes]) -> str:
    code = result.returncode
    shown = code if -256 < code < 256 else hex(code & 0xFFFFFFFF)  # Windows NTSTATUS, e.g. 0xc0000005
    reason = f"The extraction process crashed (exit {shown})"
    if code in (-9, 137):
        reason += ", killed by the system (possibly out of memory)"
    tail = _stderr_tail(result.stderr)
    return f"{reason}: {tail}" if tail else reason


def _stderr_tail(stderr: bytes | None) -> str:
    lines = [ln.strip() for ln in (stderr or b"").decode("utf-8", "replace").splitlines() if ln.strip()]
    return scrub_detail(lines[-1], 200) if lines else ""


def _tmp_sibling(path: Path) -> Path:
    return path.with_suffix(path.suffix + ".tmp")  # what NormalizedDocument.save() writes first


def _remove_quietly(path: Path) -> None:
    try:
        path.unlink(missing_ok=True)
    except OSError:
        pass


# ---- diagnostics ---------------------------------------------------------------------------


def _package_version(distribution: str) -> str | None:
    try:
        return importlib.metadata.version(distribution)
    except importlib.metadata.PackageNotFoundError:
        return None


def _resolve_tesseract(settings: DocIntelSettings, env: dict[str, str]) -> tuple[str | None, str | None]:
    """(resolved binary or None, the explicitly configured value or None): the same
    order document-extractor uses (setting / TESSERACT_CMD, PATH, Windows default)."""
    configured = (settings.tesseract_cmd or "").strip() or (env.get("TESSERACT_CMD") or "").strip() or None
    search_path = env.get("PATH")
    if configured:
        found = configured if os.path.isfile(configured) else shutil.which(configured, path=search_path)
        return found, configured
    found = shutil.which("tesseract", path=search_path)
    if found is None and os.name == "nt" and os.path.isfile(_WINDOWS_TESSERACT):
        found = _WINDOWS_TESSERACT
    return found, None


def _probe(command: list[str], env: dict[str, str]) -> str:
    kwargs: dict[str, Any] = {}
    if os.name == "nt":
        kwargs["creationflags"] = subprocess.CREATE_NO_WINDOW
    result = subprocess.run(  # noqa: S603 - fixed argv
        command, capture_output=True, timeout=_PROBE_TIMEOUT_SECONDS, env=env, stdin=subprocess.DEVNULL,
        check=False, **kwargs,
    )
    # Tesseract 4 prints to stderr, 5 to stdout.
    return (result.stdout or b"").decode("utf-8", "replace") + "\n" + (result.stderr or b"").decode("utf-8", "replace")


def _tesseract_version(path: str, env: dict[str, str]) -> tuple[str | None, str | None]:
    try:
        output = _probe([path, "--version"], env)
    except subprocess.TimeoutExpired:
        return None, "tesseract --version timed out"
    except OSError as exc:
        return None, f"Tesseract could not be started ({type(exc).__name__})"
    match = re.search(r"tesseract\s+v?(\d+(?:\.\d+)*)", output, re.IGNORECASE)
    if match is None:
        return None, "Could not read the Tesseract version"
    return match.group(1), None


def _tesseract_languages(path: str, env: dict[str, str]) -> tuple[list[str], str | None]:
    try:
        output = _probe([path, "--list-langs"], env)
    except subprocess.TimeoutExpired:
        return [], "tesseract --list-langs timed out"
    except OSError as exc:
        return [], f"Tesseract could not be started ({type(exc).__name__})"
    languages: set[str] = set()
    listing = False
    for line in output.splitlines():
        text = line.strip()
        if text.lower().startswith("list of available languages"):
            listing = True
        elif listing and re.fullmatch(r"[A-Za-z0-9_+.-]+", text):
            languages.add(text)
    if not listing:
        return [], "Could not list the Tesseract languages (is the tessdata directory set up?)"
    return sorted(languages), None


def _write_smoke_docx(path: Path) -> None:
    import docx  # python-docx, a document-extractor dependency

    document = docx.Document()
    document.add_heading("AIVA extraction smoke test", level=1)
    document.add_paragraph(f"Marker: {_SMOKE_MARKER}")
    document.add_paragraph(_SMOKE_ARABIC)
    document.save(str(path))
