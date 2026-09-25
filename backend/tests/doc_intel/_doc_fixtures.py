"""Documents generated in memory for the doc-intel storage and extraction tests.

No binary fixture is committed: every PDF is built with reportlab and every DOCX
with python-docx. Nothing here touches a database or the network.
"""
from __future__ import annotations

import io
import os
import shutil
import subprocess
import zipfile
from pathlib import Path
from typing import Any

from fastapi import UploadFile

from backend.doc_intel.settings import DocIntelSettings

PDF_MEDIA_TYPE = "application/pdf"
DOCX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

# Real Arabic with several lam-alef words ("لا", "خلال", "الاستفسار").
ARABIC_SENTENCE = "لا توجد رسوم إضافية خلال فترة الاستفسار عن الخدمة"

CFB_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
_WINDOWS_TESSERACT = r"C:\Program Files\Tesseract-OCR\tesseract.exe"


def make_settings(tmp_path: Path, **overrides: Any) -> DocIntelSettings:
    """Settings that ignore every .env file: storage under ``tmp_path``, short timeouts."""
    values: dict[str, Any] = {
        "storage_dir": str(tmp_path / "doc_intel"),
        "extraction_timeout_seconds": 120,
    }
    values.update(overrides)
    return DocIntelSettings(_env_file=None, **values)


def make_upload(data: bytes, filename: str | None, *, size: int | None = None) -> UploadFile:
    return UploadFile(file=io.BytesIO(data), filename=filename, size=size)


def make_pdf(pages: list[list[str]] | None = None, *, password: str | None = None,
             font_size: int = 11) -> bytes:
    """One page per entry; the first line of a page is set large and bold (a heading)."""
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    encrypt = None
    if password:
        from reportlab.lib import pdfencrypt

        encrypt = pdfencrypt.StandardEncryption(password, ownerPassword=password + "-owner")
    buffer = io.BytesIO()
    pdf = canvas.Canvas(buffer, pagesize=A4, encrypt=encrypt)
    for lines in pages if pages is not None else [["Document heading", "First page body text."]]:
        y = 760
        for index, line in enumerate(lines):
            pdf.setFont("Helvetica-Bold" if index == 0 else "Helvetica", 20 if index == 0 else font_size)
            pdf.drawString(72, y, line)
            y -= 36 if index == 0 else font_size * 2
        pdf.showPage()
    pdf.save()
    return buffer.getvalue()


def make_docx(*, page_breaks: int = 0) -> bytes:
    """Title, two sections, a table and an Arabic paragraph (see test_extraction for the shape)."""
    import docx

    document = docx.Document()
    document.add_heading("Customer Guide", level=1)
    document.add_paragraph("This guide explains card services for customers.")
    document.add_heading("Card Services", level=2)
    document.add_paragraph("Cards can be blocked from the mobile application.")
    document.add_paragraph(ARABIC_SENTENCE)
    table = document.add_table(rows=3, cols=2)
    for row, (service, fee) in enumerate((("Service", "Fee"), ("Replacement card", "50 EGP"), ("PIN reset", "Free"))):
        table.cell(row, 0).text = service
        table.cell(row, 1).text = fee
    document.add_heading("Fees", level=2)
    document.add_paragraph("Fees are charged monthly.")
    for index in range(page_breaks):
        document.add_page_break()
        document.add_paragraph(f"Appendix page {index + 2}.")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def rewrite_zip(data: bytes, *, add: dict[str, bytes] | None = None,
                replace: dict[str, bytes] | None = None, drop: tuple[str, ...] = ()) -> bytes:
    """Copy a ZIP/DOCX, adding, replacing or dropping members."""
    source = zipfile.ZipFile(io.BytesIO(data))
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as target:
        for info in source.infolist():
            if info.filename in drop:
                continue
            content = (replace or {}).get(info.filename, source.read(info.filename))
            target.writestr(info.filename, content)
        for name, content in (add or {}).items():
            target.writestr(name, content)
    return buffer.getvalue()


def make_zip(members: dict[str, bytes], *, compression: int = zipfile.ZIP_DEFLATED) -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", compression) as archive:
        for name, content in members.items():
            archive.writestr(name, content)
    return buffer.getvalue()


def make_cfb(stream_names: tuple[str, ...]) -> bytes:
    """Just enough of an OLE2 compound file for signature and stream-name checks."""
    names = b"".join(name.encode("utf-16-le") + b"\x00\x00" for name in stream_names)
    return CFB_MAGIC + b"\x00" * 504 + names + b"\x00" * 1024


def find_tesseract() -> str | None:
    """The Tesseract binary the extractor would use, or None."""
    for candidate in (os.environ.get("TESSERACT_CMD", "").strip(), shutil.which("tesseract") or "",
                      _WINDOWS_TESSERACT if os.name == "nt" else ""):
        if candidate and os.path.isfile(candidate):
            return candidate
    return None


def tesseract_languages(cmd: str | None) -> set[str]:
    if not cmd:
        return set()
    try:
        result = subprocess.run([cmd, "--list-langs"], capture_output=True, timeout=10, check=False)
    except (OSError, subprocess.SubprocessError):
        return set()
    text = (result.stdout or b"").decode("utf-8", "replace") + (result.stderr or b"").decode("utf-8", "replace")
    return {line.strip() for line in text.splitlines()[1:] if line.strip() and " " not in line.strip()}
