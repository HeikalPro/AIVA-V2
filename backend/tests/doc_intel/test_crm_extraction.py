"""CRM extraction of a SharePoint file (backend.doc_intel.crm_extraction + the worker's "crm" job).

The real-extraction tests start the actual child process (document-extractor, then
crm-document-ingestion's pattern extractors and validator) on DOCX files generated in
memory. Stage attribution inside the child is exercised in-process through the worker's
``_main_crm``. All names, e-mail addresses and phone numbers are fictitious.
No database, no network.
"""
from __future__ import annotations

import hashlib
import io
import json
import subprocess
import sys
from pathlib import Path

import pytest

from backend.doc_intel import crm_extraction, extraction, extraction_worker
from backend.doc_intel.crm_extraction import (
    INTELLIGENCE_NOT_SUPPORTED_REASON,
    CrmExtractionFailed,
    CrmExtractionResult,
    crm_extraction_available,
    run_crm_extraction,
)
from backend.doc_intel.extraction import run_smoke_test
from backend.doc_intel.normalized import NormalizedDocument

try:  # the test directory may or may not be a package
    from . import _doc_fixtures as fx
except ImportError:
    import _doc_fixtures as fx

ROOT = Path(__file__).resolve().parents[3]
DOCX = fx.DOCX_MEDIA_TYPE
RLM = "\u200f"

PROFILE_LINES = [
    "Company: Northwind Traders Ltd",
    f"Contact Name: {RLM}Maria Anders",  # a stray bidi mark, as Word leaves in mixed-script text
    "Email: maria.anders@northwind.example.com",
    "Phone: +44 20 7946 0958",
]
PROFILE_TABLE = [("Job Title", "Sales Director"), ("Website", "www.northwind.example.com")]
ARABIC_LINES = [
    "اسم الشركة: شركة النيل للتجارة",
    "الاسم: أحمد علي",
    "المسمى الوظيفي: مدير المبيعات",
    "البريد الإلكتروني: ahmed.ali@nile-trading.example.com",
    "رقم الهاتف: ٠١٠٠١٢٣٤٥٦٧",  # Arabic-Indic digits
]
SUPPORT_SCHEMA = {"schemas": [{
    "name": "support_case",
    "description": "Test-only client schema.",
    "fields": [
        {"name": "case_number", "type": "string", "required": True, "aliases": ["Case No"]},
        {"name": "priority", "type": "string", "aliases": ["Priority"]},
    ],
}]}


def make_docx(lines: list[str], table: list[tuple[str, str]] = ()) -> bytes:
    import docx

    document = docx.Document()
    document.add_heading("Customer Profile", level=1)
    for line in lines:
        document.add_paragraph(line)
    if table:
        grid = document.add_table(rows=len(table), cols=2)
        for row, (label, value) in enumerate(table):
            grid.cell(row, 0).text = label
            grid.cell(row, 1).text = value
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


@pytest.fixture
def settings(tmp_path):
    return fx.make_settings(tmp_path)


def run(tmp_path: Path, settings, data: bytes, *, name: str = "profile.docx", media_type: str = DOCX,
        use_intelligence: bool = False) -> CrmExtractionResult:
    source = tmp_path / f"original{Path(name).suffix}"
    source.write_bytes(data)
    return run_crm_extraction(
        source, filename=name, media_type=media_type, sha256=hashlib.sha256(data).hexdigest(),
        work_dir=tmp_path / "work", settings=settings, use_intelligence=use_intelligence,
    )


def fail(tmp_path: Path, settings, data: bytes, **kwargs) -> CrmExtractionFailed:
    with pytest.raises(CrmExtractionFailed) as info:
        run(tmp_path, settings, data, **kwargs)
    return info.value


def entity(result: CrmExtractionResult, entity_type: str) -> dict:
    matches = [e for e in result.entities if e["_meta"]["entity_type"] == entity_type]
    assert len(matches) == 1, result.entities
    return matches[0]


def forbid_children(monkeypatch) -> None:
    def refuse(*args, **kwargs):
        raise AssertionError("no extraction child may be started here")

    monkeypatch.setattr(extraction, "_spawn", refuse)


# ---- availability ----------------------------------------------------------------------------


def test_crm_extraction_available_in_this_environment():
    assert crm_extraction_available() == (True, None)


def test_a_missing_crm_library_fails_the_intelligence_stage(tmp_path, settings, monkeypatch):
    monkeypatch.setitem(sys.modules, "crm_ingestion", None)  # makes the import fail
    ok, reason = crm_extraction_available()
    assert ok is False and "crm-document-ingestion is not installed" in reason
    forbid_children(monkeypatch)
    failure = fail(tmp_path, settings, make_docx(PROFILE_LINES))
    assert (failure.stage, failure.code) == ("intelligence", "unavailable")
    assert "crm-document-ingestion" in failure.suggested_action


def test_a_missing_extractor_fails_the_extraction_stage(tmp_path, settings, monkeypatch):
    monkeypatch.setitem(sys.modules, "document_extractor", None)
    assert crm_extraction_available()[0] is False
    forbid_children(monkeypatch)
    failure = fail(tmp_path, settings, make_docx(PROFILE_LINES))
    assert (failure.stage, failure.code) == ("extraction", "unavailable")


# ---- real extractions in the child process ------------------------------------------------------


def test_profile_docx_becomes_organization_and_contact(tmp_path, settings):
    data = make_docx(PROFILE_LINES, PROFILE_TABLE)
    result = run(tmp_path, settings, data)

    assert isinstance(result.normalized, NormalizedDocument)
    assert result.normalized.filename == "profile.docx" and result.normalized.media_type == DOCX
    assert result.normalized.extractor["name"] == "document-extractor"
    assert any(b.text == "Company: Northwind Traders Ltd" for b in result.normalized.blocks)

    assert result.valid is True and result.issues == []
    org = entity(result, "organization")
    assert org["name"] == "Northwind Traders Ltd"
    assert org["website"] == "https://www.northwind.example.com"
    contact = entity(result, "contact")
    assert contact["full_name"] == "Maria Anders"  # the bidi mark is gone
    assert contact["job_title"] == "Sales Director"
    assert contact["email"] == "maria.anders@northwind.example.com"
    assert contact["phone"] == "+442079460958"
    assert len(result.entities) == 2

    # provenance: each value points back at the text it was read from
    meta = contact["_meta"]
    assert meta["entity_type"] == "contact" and 0 < meta["confidence"] <= 1
    email = meta["fields"]["email"]
    assert email["extractor"] == "pattern"
    assert email["provenance"][0]["source_text"] == "Email: maria.anders@northwind.example.com"
    assert email["provenance"][0]["block_id"]
    assert meta["fields"]["job_title"]["provenance"][0]["source_text"] == "Job Title: Sales Director"  # the table

    # the full result (stored as result_json) and its source metadata
    assert set(result.result) == {"extraction", "validation"}
    source = result.result["extraction"]["source"]
    assert (source["source_system"], source["filename"], source["mime_type"]) == ("sharepoint", "profile.docx", DOCX)
    assert source["extra"]["sha256"] == hashlib.sha256(data).hexdigest() and source["size"] == len(data)
    assert result.result["validation"]["valid"] is True
    assert result.result["extraction"]["extractors"] == ["pattern"]
    json.dumps(result.result)  # plain JSON

    metrics = result.metrics
    assert metrics["entity_count"] == 2 and metrics["entity_types"] == {"organization": 1, "contact": 1}
    assert metrics["extractors"] == ["pattern"] and metrics["use_intelligence"] is False
    assert metrics["min_confidence"] == 0.5 and metrics["valid"] is True
    assert metrics["schemas"]["organization"]["display_field"] == "name"
    assert metrics["schemas"]["contact"]["display_field"] == "full_name"
    assert metrics["schemas"]["contact"]["field_types"]["email"] == "email"
    assert metrics["crm_ingestion_version"] and metrics["document_extractor_version"]
    assert set(metrics["seconds"]) == {"extraction", "intelligence", "entities", "total"}
    assert all(isinstance(v, float) for v in metrics["seconds"].values())

    # only the job files are left in the caller's work_dir
    assert sorted(p.name for p in (tmp_path / "work").iterdir()) == ["crm_job.json", "crm_progress.txt",
                                                                    "crm_result.json"]


def test_arabic_labels_and_digits(tmp_path, settings):
    result = run(tmp_path, settings, make_docx(ARABIC_LINES), name="ملف العميل.docx")
    assert result.valid is True, result.issues
    assert entity(result, "organization")["name"] == "شركة النيل للتجارة"
    contact = entity(result, "contact")
    assert contact["full_name"] == "أحمد علي"
    assert contact["job_title"] == "مدير المبيعات"
    assert contact["email"] == "ahmed.ali@nile-trading.example.com"
    assert contact["phone"] == "01001234567"
    assert result.normalized.filename == "ملف العميل.docx"


def test_validation_issues_make_the_result_invalid(tmp_path, settings):
    result = run(tmp_path, settings, make_docx(["Email: sara.ahmed@acme.example.com"]))
    assert result.valid is False
    assert {"entity_type": "contact", "entity_index": 0, "field": "full_name", "severity": "error",
            "code": "missing_required_field", "message": "required field 'full_name' is missing"} in result.issues
    assert result.metrics["valid"] is False and result.metrics["error_count"] >= 1
    assert entity(result, "contact")["email"] == "sara.ahmed@acme.example.com"


def test_a_client_schema_file_is_loaded(tmp_path):
    schema = tmp_path / "client.json"
    schema.write_text(json.dumps(SUPPORT_SCHEMA), encoding="utf-8")
    settings = fx.make_settings(tmp_path, crm_schema_paths=str(schema), crm_min_confidence=0.9)
    result = run(tmp_path, settings, make_docx(["Case No: CS-1042", "Priority: High"]))
    case = entity(result, "support_case")
    assert (case["case_number"], case["priority"]) == ("CS-1042", "High")
    assert "support_case" in result.metrics["schemas"] and result.metrics["min_confidence"] == 0.9
    # 0.85 label confidence is below the stricter threshold: a warning, not an error
    assert any(i["code"] == "low_confidence" and i["severity"] == "warning" for i in result.issues)
    assert result.valid is True


def test_schema_paths_resolve_relative_to_the_repository_root(tmp_path):
    settings = fx.make_settings(tmp_path, crm_schema_paths=f"schemas/client.json, {tmp_path / 'abs.json'}")
    assert crm_extraction.schema_paths(settings) == [ROOT / "schemas" / "client.json", tmp_path / "abs.json"]


# ---- stage-specific failures ------------------------------------------------------------------------


def test_a_corrupt_file_fails_at_the_extraction_stage(tmp_path, settings):
    failure = fail(tmp_path, settings, b"PK\x03\x04" + bytes(range(256)) * 20)
    assert failure.stage == "extraction"
    assert failure.code in {"unsupported", "parse_error", "resource_limit"}
    assert str(tmp_path) not in failure.reason and ":\\" not in failure.reason


def test_a_password_protected_pdf_fails_at_the_extraction_stage(tmp_path, settings):
    failure = fail(tmp_path, settings, fx.make_pdf(password="fake-pass"), name="locked.pdf",
                   media_type=fx.PDF_MEDIA_TYPE)
    assert (failure.stage, failure.code) == ("extraction", "password_protected")


@pytest.mark.parametrize(
    "content, code, fragment",
    [("{not json", "schema_error", "is invalid"),
     (json.dumps({"name": "organization", "fields": [{"name": "x"}]}), "schema_error", "already registered"),
     (None, "schema_error", "could not be read")],
)
def test_a_bad_schema_file_fails_at_the_intelligence_stage(tmp_path, content, code, fragment):
    schema = tmp_path / "private" / "client.json"
    if content is not None:
        schema.parent.mkdir()
        schema.write_text(content, encoding="utf-8")
    settings = fx.make_settings(tmp_path, crm_schema_paths=str(schema))
    failure = fail(tmp_path, settings, make_docx(PROFILE_LINES))
    assert (failure.stage, failure.code) == ("intelligence", code)
    assert "client.json" in failure.reason and fragment in failure.reason
    assert str(schema.parent) not in failure.reason
    assert "DOC_INTEL_CRM_SCHEMA_PATHS" in failure.suggested_action


def test_missing_or_unsupported_sources_fail_before_a_child(tmp_path, settings, monkeypatch):
    forbid_children(monkeypatch)
    with pytest.raises(CrmExtractionFailed) as info:
        run_crm_extraction(tmp_path / "gone.docx", filename="gone.docx", media_type=DOCX, sha256="0" * 64,
                           work_dir=tmp_path / "work", settings=settings)
    assert (info.value.stage, info.value.code) == ("extraction", "source_missing")
    failure = fail(tmp_path, settings, b"plain text", name="notes.txt", media_type="text/plain")
    assert (failure.stage, failure.code) == ("extraction", "unsupported")


# ---- stage attribution inside the child (in-process) ---------------------------------------------


def crm_job_file(tmp_path: Path, data: bytes, **crm) -> dict:
    source = tmp_path / "original.docx"
    source.write_bytes(data)
    job = extraction._job_base(source, filename="profile.docx", media_type=DOCX,
                               sha256=hashlib.sha256(data).hexdigest(), settings=fx.make_settings(tmp_path),
                               timeout=120.0)
    job.update({
        "kind": "crm",
        "output_path": str(tmp_path / "crm_result.json"),
        "error_path": str(tmp_path / "crm_error.json"),
        "progress_path": str(tmp_path / "crm_progress.txt"),
        "crm": {"schema_paths": [], "min_confidence": 0.5, "use_intelligence": False, **crm},
    })
    return job


def child_error(tmp_path: Path) -> dict:
    return json.loads((tmp_path / "crm_error.json").read_text(encoding="utf-8"))["error"]


def test_an_extractor_failure_is_the_intelligence_stage(tmp_path, monkeypatch):
    from crm_ingestion.crm import PatternExtractor

    def broken(self, document, *, schemas):
        raise RuntimeError("pattern bug")

    monkeypatch.setattr(PatternExtractor, "extract", broken)
    job = crm_job_file(tmp_path, make_docx(PROFILE_LINES))
    assert extraction_worker._main_crm(job, Path(job["error_path"])) == extraction_worker.EXIT_FAILED
    error = child_error(tmp_path)
    assert (error["stage"], error["code"]) == ("intelligence", "entity_extraction_failed")
    assert "pattern bug" in error["reason"]


def test_a_failure_after_the_extractors_is_the_entities_stage(tmp_path, monkeypatch):
    from crm_ingestion.crm.validators.validator import EntityValidator

    def broken(self, entities):
        raise RuntimeError("validator bug")

    monkeypatch.setattr(EntityValidator, "validate", broken)
    job = crm_job_file(tmp_path, make_docx(PROFILE_LINES))
    assert extraction_worker._main_crm(job, Path(job["error_path"])) == extraction_worker.EXIT_CRASH
    error = child_error(tmp_path)
    assert (error["stage"], error["code"]) == ("entities", "crash")
    assert (tmp_path / "crm_progress.txt").read_text(encoding="utf-8") == "entities"
    assert not (tmp_path / "crm_result.json").exists()


def test_the_worker_itself_refuses_llm_intelligence(tmp_path):
    job = crm_job_file(tmp_path, make_docx(PROFILE_LINES), use_intelligence=True)
    assert extraction_worker._main_crm(job, Path(job["error_path"])) == extraction_worker.EXIT_FAILED
    error = child_error(tmp_path)
    assert (error["stage"], error["code"]) == ("intelligence", "not_supported")


def test_flow1_error_files_keep_their_shape(tmp_path):
    path = tmp_path / "error.json"
    extraction_worker._write_error(path, "parse_error", "The file could not be parsed", None)
    assert json.loads(path.read_text(encoding="utf-8")) == {
        "error": {"code": "parse_error", "reason": "The file could not be parsed", "suggested_action": None}}


def test_the_job_kind_is_validated(tmp_path):
    job = crm_job_file(tmp_path, make_docx(PROFILE_LINES))
    for broken in ({**job, "kind": "shell"}, {**job, "crm": None}):
        path = tmp_path / "job.json"
        path.write_text(json.dumps(broken), encoding="utf-8")
        with pytest.raises((ValueError, KeyError)):
            extraction_worker.load_job(path)


# ---- the parent side: intelligence, the slot, timeouts and crashes --------------------------------


def test_use_intelligence_is_refused_before_any_child(tmp_path, settings, monkeypatch):
    forbid_children(monkeypatch)
    failure = fail(tmp_path, settings, make_docx(PROFILE_LINES), use_intelligence=True)
    assert (failure.stage, failure.code) == ("intelligence", "not_supported")
    assert failure.reason == INTELLIGENCE_NOT_SUPPORTED_REASON == "LLM-based CRM intelligence is not enabled in this version"
    assert failure.suggested_action


def test_the_extraction_slot_is_shared_with_flow_1(tmp_path, settings, monkeypatch):
    seen: dict[str, object] = {}

    def fake_child(*args, **kwargs):
        seen["slot_taken"] = not extraction._EXTRACTION_SLOT.acquire(blocking=False)
        seen["smoke"] = run_smoke_test(settings, timeout_seconds=5)  # Flow 1's health check skips, busy
        return CrmExtractionResult(
            normalized=NormalizedDocument(filename="x.docx", media_type=DOCX, sha256="0" * 64),
            entities=[], result={}, valid=True,
        )

    monkeypatch.setattr(crm_extraction, "_run_crm_child", fake_child)
    run(tmp_path, settings, make_docx(PROFILE_LINES))
    assert seen["slot_taken"] is True
    assert seen["smoke"]["busy"] is True
    assert extraction._EXTRACTION_SLOT.acquire(blocking=False)  # released afterwards
    extraction._EXTRACTION_SLOT.release()


def capture_job(command) -> dict:
    return json.loads(Path(command[-1]).read_text(encoding="utf-8"))


def test_the_job_carries_explicit_options_and_the_child_gets_a_scrubbed_env(tmp_path, monkeypatch):
    monkeypatch.setenv("MICROSOFT_CLIENT_SECRET", "fake-env-secret")
    monkeypatch.setenv("MICROSOFT_TENANT_ID", "fake-env-tenant")
    monkeypatch.setenv("CRM_SCHEMA_PATHS", '["evil.json"]')
    monkeypatch.setenv("DOC_INTEL_SECRETS_KEY", "fake-fernet-key")
    schema = tmp_path / "client.json"
    settings = fx.make_settings(tmp_path, crm_schema_paths=str(schema), crm_min_confidence=0.7)
    seen: dict[str, object] = {}

    def capture(command, *, cwd, env, timeout):
        seen.update(job=capture_job(command), command=command, cwd=cwd, env=env, timeout=timeout)
        return subprocess.CompletedProcess(command, 3, b"", b"")

    monkeypatch.setattr(extraction, "_spawn", capture)
    fail(tmp_path, settings, make_docx(PROFILE_LINES))
    job = seen["job"]
    assert job["kind"] == "crm" and job["version"] == extraction_worker.JOB_VERSION
    assert job["crm"] == {"schema_paths": [str(schema)], "min_confidence": 0.7, "use_intelligence": False}
    assert job["mode"] == "balanced" and job["ocr_languages"] == ["ara", "eng"] and job["max_pages"] == 300
    assert Path(job["output_path"]).parent == (tmp_path / "work").resolve()
    assert seen["command"][:4] == [sys.executable, "-m", "backend.doc_intel.extraction_worker", "--job"]
    assert Path(seen["cwd"]) == ROOT and seen["timeout"] == 120.0
    upper = {k.upper() for k in seen["env"]}
    assert not upper & {"MICROSOFT_CLIENT_SECRET", "MICROSOFT_TENANT_ID", "CRM_SCHEMA_PATHS", "DOC_INTEL_SECRETS_KEY"}


def test_a_timeout_names_the_stage_the_child_was_in(tmp_path, settings, monkeypatch):
    def slow(command, *, cwd, env, timeout):
        Path(capture_job(command)["progress_path"]).write_text("intelligence", encoding="utf-8")
        raise subprocess.TimeoutExpired(command, timeout)

    monkeypatch.setattr(extraction, "_spawn", slow)
    failure = fail(tmp_path, settings, make_docx(PROFILE_LINES))
    assert (failure.stage, failure.code) == ("intelligence", "timeout")
    assert failure.reason.startswith("Extraction timed out after 120 s")


def test_a_crash_without_an_error_file_names_the_stage_and_the_stderr_tail(tmp_path, settings, monkeypatch):
    def crash(command, *, cwd, env, timeout):
        Path(capture_job(command)["progress_path"]).write_text("entities", encoding="utf-8")
        return subprocess.CompletedProcess(command, 3, b"", b"Traceback...\nRuntimeError: boom in C:\\srv\\x.py\n")

    monkeypatch.setattr(extraction, "_spawn", crash)
    failure = fail(tmp_path, settings, make_docx(PROFILE_LINES))
    assert (failure.stage, failure.code) == ("entities", "crash")
    assert failure.reason == "The extraction process crashed (exit 3): RuntimeError: boom in <path>"


def test_the_child_error_file_is_mapped_with_its_stage(tmp_path, settings, monkeypatch):
    def reported(command, *, cwd, env, timeout):
        Path(capture_job(command)["error_path"]).write_text(json.dumps({"error": {
            "code": "schema_error", "reason": "The CRM schema file c.json is invalid at C:\\srv\\c.json",
            "suggested_action": "Fix it", "stage": "intelligence"}}), encoding="utf-8")
        return subprocess.CompletedProcess(command, 1, b"", b"")

    monkeypatch.setattr(extraction, "_spawn", reported)
    failure = fail(tmp_path, settings, make_docx(PROFILE_LINES))
    assert (failure.stage, failure.code, failure.suggested_action) == ("intelligence", "schema_error", "Fix it")
    assert "C:\\srv" not in failure.reason and "<path>" in failure.reason


def test_a_stale_result_from_an_earlier_attempt_is_never_used(tmp_path, settings, monkeypatch):
    work = tmp_path / "work"
    work.mkdir()
    (work / "crm_result.json").write_text(json.dumps({"normalized": {}, "entities": [], "result": {}}),
                                          encoding="utf-8")

    def exits_cleanly_without_a_result(command, *, cwd, env, timeout):
        return subprocess.CompletedProcess(command, 0, b"", b"")

    monkeypatch.setattr(extraction, "_spawn", exits_cleanly_without_a_result)
    failure = fail(tmp_path, settings, make_docx(PROFILE_LINES))
    assert (failure.stage, failure.code) == ("entities", "crash")
    assert "without a usable result" in failure.reason


def test_a_real_timeout_kills_the_crm_child(tmp_path, settings):
    source = tmp_path / "original.docx"
    source.write_bytes(make_docx(PROFILE_LINES))
    (tmp_path / "work").mkdir()
    with pytest.raises(CrmExtractionFailed) as info:
        crm_extraction._run_crm_child(source, filename="p.docx", media_type=DOCX, sha256="0" * 64,
                                      work_dir=tmp_path / "work", settings=settings, timeout=0.05)
    assert (info.value.stage, info.value.code) == ("extraction", "timeout")
    assert not (tmp_path / "work" / "crm_result.json").exists()
