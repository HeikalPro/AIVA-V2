"""The whole chain with no Microsoft calls: fake Graph (httpx.MockTransport) -> DocumentDownloader
-> IngestionPipeline (real document-extractor or a fake) -> CRMExtractionService, plus the
`CRMIngestionApp` façade and the `crm-ingest` CLI."""

from __future__ import annotations

import io
import json
import subprocess
import sys
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

import httpx
import pytest
from connectors.fakes import SP_DRIVE_ID, SP_ITEM_ID, graph_settings, make_client, sharepoint_item
from crm.fakes import business_letter
from ingestion.fakes import FakeExtractor

from crm_ingestion import (
    ConfigurationError,
    CRMExtractionService,
    CRMIngestionApp,
    CRMProcessingResult,
    DocumentDownloader,
    DriveItemReference,
    IngestionPipeline,
    Settings,
)
from crm_ingestion.cli import main
from crm_ingestion.config import CRMSettings, ExtractionSettings
from crm_ingestion.connectors.sharepoint import encode_sharing_url

NOW = datetime(2026, 9, 24, 12, 0, tzinfo=UTC)
DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
FILENAME = "Northwind Traders - Profile.docx"
DOWNLOAD_URL = "https://contoso.sharepoint.com/sites/Sales/_layouts/15/download.aspx?UniqueId=nw&tempauth=t"
WEB_URL = (
    "https://contoso.sharepoint.com/sites/Sales/Shared%20Documents/Northwind%20Traders%20-%20Profile.docx"
)
ETAG = '"{11111111-2222-3333-4444-555555555555},7"'
SHARING_URL = "https://contoso.sharepoint.com/:w:/s/Sales/EabcDEF123?e=xyz"


def northwind_docx() -> bytes:
    """A real DOCX: organisation name, labelled contact lines and a 2-column label/value table."""
    import docx

    d = docx.Document()
    d.add_heading("Northwind Traders Ltd", level=1)
    d.add_paragraph("Company: Northwind Traders Ltd")
    d.add_paragraph("Email: sales@northwind.example.com")
    d.add_paragraph("Phone: +44 20 7946 0958")
    rows = [
        ("Contact Name", "Maria Anders"),
        ("Job Title", "Sales Director"),
        ("Website", "www.northwind.example.com"),
    ]
    table = d.add_table(rows=len(rows), cols=2)
    for r, (label, value) in enumerate(rows):
        table.cell(r, 0).text = label
        table.cell(r, 1).text = value
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


class FakeGraph:
    """Serves one driveItem (by ids or sharing URL) and its content; records every request."""

    def __init__(self, content: bytes) -> None:
        self.content = content
        self.requests: list[httpx.Request] = []

    def item(self) -> dict[str, Any]:
        data = sharepoint_item(
            name=FILENAME,
            size=len(self.content),
            webUrl=WEB_URL,
            eTag=ETAG,
            file={"mimeType": DOCX_MIME, "hashes": {"quickXorHash": "cXV4"}},
        )
        data["@microsoft.graph.downloadUrl"] = DOWNLOAD_URL
        return data

    def __call__(self, request: httpx.Request) -> httpx.Response:
        self.requests.append(request)
        path = request.url.path
        if request.url.host == "graph.microsoft.com":
            if path == f"/v1.0/drives/{SP_DRIVE_ID}/items/{SP_ITEM_ID}":
                return httpx.Response(200, json=self.item())
            if path == f"/v1.0/shares/{encode_sharing_url(SHARING_URL)}/driveItem":
                return httpx.Response(200, json=self.item())
            return httpx.Response(404, json={"error": {"code": "itemNotFound", "message": path}})
        if str(request.url) == DOWNLOAD_URL:
            return httpx.Response(200, content=self.content)
        return httpx.Response(500)

    def downloader(self) -> DocumentDownloader:
        return DocumentDownloader(make_client(self, settings=graph_settings()), clock=lambda: NOW)


def fast_settings(**crm: Any) -> Settings:
    return Settings(extraction=ExtractionSettings(mode="fast"), crm=CRMSettings(**crm))


def real_pipeline() -> IngestionPipeline:
    return IngestionPipeline(settings=fast_settings())


def entity(result: CRMProcessingResult, entity_type: str) -> dict[str, Any]:
    matches = [e for e in result.to_crm_json() if e["_meta"]["entity_type"] == entity_type]
    assert matches, f"no {entity_type} entity in {result.to_crm_json()}"
    return matches[0]


def assert_northwind(result: CRMProcessingResult) -> None:
    assert result.valid, result.validation
    org = entity(result, "organization")
    assert org["name"] == "Northwind Traders Ltd"
    assert org["website"] == "https://www.northwind.example.com"
    contact = entity(result, "contact")
    assert contact["full_name"] == "Maria Anders"
    assert contact["job_title"] == "Sales Director"
    assert contact["email"] == "sales@northwind.example.com"
    assert contact["phone"] == "+442079460958"
    # provenance: every value points back at the block and text it came from
    email_prov = contact["_meta"]["fields"]["email"]["provenance"]
    assert email_prov[0]["source_text"] == "Email: sales@northwind.example.com"
    assert email_prov[0]["page"] == 1 and email_prov[0]["block_id"]
    name_prov = org["_meta"]["fields"]["name"]["provenance"]
    assert "Northwind Traders Ltd" in name_prov[0]["source_text"]
    assert org["_meta"]["source_document_id"] == result.extraction.document_id


def assert_sharepoint_source(source: dict[str, Any] | None) -> None:
    assert source is not None
    assert source["source_system"] == "sharepoint"
    assert source["filename"] == FILENAME
    assert (source["drive_id"], source["item_id"]) == (SP_DRIVE_ID, SP_ITEM_ID)
    assert source["source_uri"] == WEB_URL
    assert source["etag"] == ETAG
    assert source["mime_type"] == DOCX_MIME
    assert source["retrieved_at"] == "2026-09-24T12:00:00Z"


# ---- layer by layer ----------------------------------------------------------------


def test_sharepoint_docx_to_crm_json_with_real_extractor() -> None:
    graph = FakeGraph(northwind_docx())
    downloaded = graph.downloader().download(DriveItemReference.from_ids(SP_DRIVE_ID, SP_ITEM_ID))
    ingested = real_pipeline().ingest(downloaded)
    assert ingested.document.media_type == DOCX_MIME
    assert ingested.warnings == []

    result = CRMExtractionService(settings=CRMSettings()).process(ingested)

    assert_northwind(result)
    assert_sharepoint_source(result.extraction.source)
    assert ingested.document.metadata["source"]["etag"] == ETAG  # also travels on the Document
    # the pre-authenticated download URL got no bearer token; Graph metadata did
    by_host = {r.url.host: r for r in graph.requests}
    assert "authorization" not in by_host["contoso.sharepoint.com"].headers
    assert by_host["graph.microsoft.com"].headers["authorization"] == "Bearer test-token"
    # the JSON round-trips
    assert json.loads(result.to_json())["extraction"]["source"]["item_id"] == SP_ITEM_ID


def test_sharing_url_with_fake_extractor() -> None:
    graph = FakeGraph(b"not really a docx; the extractor is fake")
    extractor = FakeExtractor(document=business_letter())
    pipeline = IngestionPipeline(extractor)
    ingested = pipeline.ingest_from(graph.downloader(), DriveItemReference.from_sharing_url(SHARING_URL))

    result = CRMExtractionService(settings=CRMSettings()).process(ingested)

    assert extractor.calls == [(graph.content, FILENAME)]
    assert graph.requests[0].url.path == f"/v1.0/shares/{encode_sharing_url(SHARING_URL)}/driveItem"
    assert_sharepoint_source(result.extraction.source)
    assert result.extraction.source is not None
    assert result.extraction.source["extra"]["sharing_url"] == SHARING_URL
    org = entity(result, "organization")
    assert org["name"] == "ACME Trading LLC"
    prov = org["_meta"]["fields"]["name"]["provenance"][0]
    assert prov["block_id"] in {"b0", "b1"} and "ACME Trading LLC" in prov["source_text"]
    assert entity(result, "contact")["email"] == "sara.ahmed@acme-trading.example.com"


# ---- CRM_SCHEMA_PATHS ---------------------------------------------------------------

SUPPORT_SCHEMA = {
    "schemas": [
        {
            "name": "support_case",
            "description": "Test-only client schema loaded from a JSON file.",
            "fields": [
                {"name": "case_number", "type": "string", "required": True, "aliases": ["Case No"]},
                {"name": "priority", "type": "string", "aliases": ["Priority"]},
            ],
        }
    ]
}


def test_schema_paths_are_loaded_on_top_of_generic(tmp_path: Path) -> None:
    path = tmp_path / "client.json"
    path.write_text(json.dumps(SUPPORT_SCHEMA), encoding="utf-8")
    service = CRMExtractionService(settings=CRMSettings(schema_paths=[path]))
    assert service.registry.names() == ["organization", "contact", "document_reference", "support_case"]

    from crm.fakes import make_document, make_ingested, table_block

    doc = make_document([table_block("t0", [["Case No", "CS-1042"], ["Priority", "High"]])])
    result = service.process(make_ingested(doc))
    case = entity(result, "support_case")
    assert (case["case_number"], case["priority"]) == ("CS-1042", "High")


def test_schema_paths_from_env_and_explicit_registry_wins(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    path = tmp_path / "client.json"
    path.write_text(json.dumps(SUPPORT_SCHEMA), encoding="utf-8")
    monkeypatch.setenv("CRM_SCHEMA_PATHS", json.dumps([str(path)]))
    assert "support_case" in CRMExtractionService().registry
    from crm_ingestion import default_registry

    assert "support_case" not in CRMExtractionService(registry=default_registry()).registry


def test_schema_paths_duplicate_generic_name_is_rejected(tmp_path: Path) -> None:
    path = tmp_path / "dup.json"
    path.write_text(json.dumps({"name": "organization", "fields": [{"name": "x"}]}), encoding="utf-8")
    with pytest.raises(ValueError, match="already registered"):
        CRMExtractionService(settings=CRMSettings(schema_paths=[path]))


# ---- façade -------------------------------------------------------------------------


def test_app_process_drive_item_and_sharing_url() -> None:
    graph = FakeGraph(northwind_docx())
    app = CRMIngestionApp(settings=fast_settings(), pipeline=real_pipeline(), downloader=graph.downloader())
    with app:
        by_ids = app.process_drive_item(SP_DRIVE_ID, SP_ITEM_ID)
        by_url = app.process_sharing_url(SHARING_URL)
    assert_northwind(by_ids)
    assert_sharepoint_source(by_ids.extraction.source)
    assert by_url.to_crm_json()[0]["name"] == by_ids.to_crm_json()[0]["name"]


def test_app_process_file_never_builds_downloader(tmp_path: Path) -> None:
    path = tmp_path / FILENAME
    path.write_bytes(northwind_docx())
    built: list[int] = []

    def factory() -> DocumentDownloader:
        built.append(1)
        raise AssertionError("downloader must not be built for local files")

    app = CRMIngestionApp(settings=fast_settings(), downloader_factory=factory)
    result = app.process_file(path)

    assert built == []
    assert_northwind(result)
    source = result.extraction.source
    assert source is not None and source["source_system"] == "local"
    assert source["filename"] == FILENAME and source["source_uri"] == path.resolve().as_uri()


def _forbid_network(monkeypatch: pytest.MonkeyPatch) -> None:
    def refuse(self: httpx.Client, request: httpx.Request, **kwargs: Any) -> httpx.Response:
        raise AssertionError(f"unexpected HTTP request: {request.method} {request.url}")

    monkeypatch.setattr(httpx.Client, "send", refuse)


def test_app_sharepoint_without_credentials_fails_before_any_http(monkeypatch: pytest.MonkeyPatch) -> None:
    _forbid_network(monkeypatch)
    app = CRMIngestionApp(settings=fast_settings())  # default downloader, no credentials
    with pytest.raises(ConfigurationError) as info:
        app.process_sharing_url(SHARING_URL)
    for name in ("MICROSOFT_TENANT_ID", "MICROSOFT_CLIENT_ID", "MICROSOFT_CLIENT_SECRET"):
        assert name in str(info.value)
    assert isinstance(app.downloader, DocumentDownloader)  # built lazily by the failed call


# ---- CLI ----------------------------------------------------------------------------


def _docx_file(tmp_path: Path) -> Path:
    path = tmp_path / FILENAME
    path.write_bytes(northwind_docx())
    return path


def test_cli_file_full_result_to_stdout(tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    code = main(["file", str(_docx_file(tmp_path)), "--mode", "fast"])
    out = capsys.readouterr()
    assert code == 0, out.err
    data = json.loads(out.out)
    assert data["validation"]["valid"] is True
    assert data["extraction"]["source"]["filename"] == FILENAME
    names = {e["entity_type"]: e for e in data["extraction"]["entities"]}
    assert names["organization"]["fields"]["name"]["value"] == "Northwind Traders Ltd"


def test_cli_file_crm_only_to_output_file(tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    target = tmp_path / "out.json"
    code = main(["file", str(_docx_file(tmp_path)), "--mode", "fast", "--crm-only", "-o", str(target)])
    out = capsys.readouterr()
    assert code == 0, out.err
    assert out.out == ""
    assert "wrote 2 entities" in out.err and "(valid)" in out.err
    data = json.loads(target.read_text(encoding="utf-8"))
    assert [e["_meta"]["entity_type"] for e in data] == ["organization", "contact"]
    assert data[1]["email"] == "sales@northwind.example.com"


def test_cli_file_errors(tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    assert main(["file", str(tmp_path / "missing.docx")]) == 2
    assert "file not found" in capsys.readouterr().err

    path = _docx_file(tmp_path)
    failing = FakeExtractor(error=ValueError("corrupt"))

    def app_factory(settings: Settings) -> CRMIngestionApp:
        return CRMIngestionApp(settings=settings, pipeline=IngestionPipeline(failing))

    assert main(["file", str(path)], app_factory=app_factory) == 2
    err = capsys.readouterr().err
    assert err.startswith("crm-ingest: error: DocumentExtractionFailed:") and "corrupt" in err
    assert err.count("\n") == 1


def test_cli_sharepoint_with_injected_downloader(capsys: pytest.CaptureFixture[str]) -> None:
    graph = FakeGraph(northwind_docx())
    seen: list[Settings] = []

    def app_factory(settings: Settings) -> CRMIngestionApp:
        seen.append(settings)
        return CRMIngestionApp(settings=settings, downloader=graph.downloader())

    code = main(
        ["sharepoint", "--drive-id", SP_DRIVE_ID, "--item-id", SP_ITEM_ID, "--mode", "fast", "--crm-only"],
        app_factory=app_factory,
    )
    out = capsys.readouterr()
    assert code == 0, out.err
    assert seen[0].extraction.mode == "fast"
    data = json.loads(out.out)
    assert data[0]["name"] == "Northwind Traders Ltd"

    code = main(["sharepoint", "--url", SHARING_URL, "--mode", "fast"], app_factory=app_factory)
    out = capsys.readouterr()
    assert code == 0, out.err
    assert json.loads(out.out)["extraction"]["source"]["extra"]["sharing_url"] == SHARING_URL


def test_cli_sharepoint_without_credentials(
    monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str]
) -> None:
    _forbid_network(monkeypatch)
    code = main(["sharepoint", "--url", "https://contoso.sharepoint.com/:w:/s/x/abc"])
    out = capsys.readouterr()
    assert code == 2
    assert out.out == ""
    assert out.err.startswith("crm-ingest: error: ConfigurationError:")
    assert "MICROSOFT_TENANT_ID, MICROSOFT_CLIENT_ID, MICROSOFT_CLIENT_SECRET" in out.err


def test_cli_sharepoint_argument_errors(capsys: pytest.CaptureFixture[str]) -> None:
    assert main(["sharepoint", "--drive-id", "d"]) == 2
    assert "--drive-id requires --item-id" in capsys.readouterr().err
    assert main(["sharepoint", "--url", "u", "--drive-id", "d"]) == 2
    assert main(["sharepoint", "--url", "u", "--item-id", "i"]) == 2
    assert main([]) == 2


def test_cli_config_check_without_credentials(capsys: pytest.CaptureFixture[str]) -> None:
    assert main(["config", "check"]) == 0
    out = capsys.readouterr().out
    assert "MICROSOFT_CLIENT_SECRET" in out and "missing" in out
    assert "SharePoint is not configured" in out
    assert "MICROSOFT_TENANT_ID, MICROSOFT_CLIENT_ID, MICROSOFT_CLIENT_SECRET" in out
    assert "organization, contact, document_reference" in out


def test_cli_config_check_never_prints_secrets(
    monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str]
) -> None:
    monkeypatch.setenv("MICROSOFT_TENANT_ID", "tenant-guid-123")
    monkeypatch.setenv("MICROSOFT_CLIENT_ID", "client-guid-456")
    monkeypatch.setenv("MICROSOFT_CLIENT_SECRET", "s3cr3t-value")
    assert main(["config", "check"]) == 0
    out = capsys.readouterr().out
    for value in ("tenant-guid-123", "client-guid-456", "s3cr3t-value"):
        assert value not in out
    assert "SharePoint is configured" in out


def test_cli_config_check_reports_bad_schema_file_and_llm_config(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str]
) -> None:
    monkeypatch.setenv("CRM_SCHEMA_PATHS", json.dumps([str(tmp_path / "nope.json")]))
    toml = tmp_path / "de.toml"
    toml.write_text('[intelligence]\nenabled = true\nbase_url = "http://llm.local/v1"\n', encoding="utf-8")
    monkeypatch.setenv("DOCUMENT_EXTRACTOR_CONFIG", str(toml))
    assert main(["config", "check"]) == 1
    out = capsys.readouterr().out
    assert "NOT FOUND" in out and "schemas could not be loaded" in out
    assert "document text is sent to http://llm.local/v1" in out


def test_cli_schemas_list(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str]
) -> None:
    assert main(["schemas", "list"]) == 0
    out = capsys.readouterr().out
    assert out.splitlines()[0].startswith("organization")
    assert "support_case" not in out

    path = tmp_path / "client.json"
    path.write_text(json.dumps(SUPPORT_SCHEMA), encoding="utf-8")
    monkeypatch.setenv("CRM_SCHEMA_PATHS", json.dumps([str(path)]))
    assert main(["schemas", "list", "--json"]) == 0
    names = [s["name"] for s in json.loads(capsys.readouterr().out)["schemas"]]
    assert names == ["organization", "contact", "document_reference", "support_case"]


def test_python_dash_m_entry_point() -> None:
    proc = subprocess.run(
        [sys.executable, "-m", "crm_ingestion", "--version"], capture_output=True, text=True, check=False
    )
    assert proc.returncode == 0
    assert proc.stdout.strip() == "crm-ingest 0.1.0"
